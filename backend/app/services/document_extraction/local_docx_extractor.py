from pathlib import Path

from app.schemas.extraction import ExtractedDocument, ExtractedPage, ExtractedTable
from app.services.document_extraction.base import DocumentExtractor


class LocalDocxExtractor(DocumentExtractor):
    def extract(self, file_path: Path) -> ExtractedDocument:
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX extraction") from exc

        doc = docx.Document(file_path)
        full_text_parts: list[str] = []
        tables: list[ExtractedTable] = []

        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if text:
                full_text_parts.append(text)

        for table in doc.tables:
            rows_data: list[list[str]] = []
            for row in table.rows:
                row_values = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_values)
                full_text_parts.append(" | ".join(row_values))
            tables.append(ExtractedTable(page_number=1, rows=rows_data))

        section_text: list[str] = []
        for section in doc.sections:
            if section.header:
                header_lines = [p.text.strip() for p in section.header.paragraphs if p.text.strip()]
                if header_lines:
                    section_text.append("HEADER: " + " ".join(header_lines))
            if section.footer:
                footer_lines = [p.text.strip() for p in section.footer.paragraphs if p.text.strip()]
                if footer_lines:
                    section_text.append("FOOTER: " + " ".join(footer_lines))

        all_text = "\n".join(section_text + full_text_parts)
        return ExtractedDocument(
            file_name=file_path.name,
            file_type="docx",
            page_count=1,
            extraction_method="docx_parser",
            full_text=all_text,
            pages=[ExtractedPage(page_number=1, text=all_text, blocks=[])],
            tables=tables,
            detected_sections=[],
            warnings=[],
        )
